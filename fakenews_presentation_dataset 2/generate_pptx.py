"""
Veritas AI Fact-Checking System - PowerPoint Presentation Generator
Theme: Modern Red and White with native vector flowcharts, native tables, and full 12-claim coverage.
"""

import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = "/Users/amanpanwar/Desktop/fakenews_presentation_dataset 2/Veritas_Fake_News_Detection_Presentation.pptx"
CLAIMS_CSV = "/Users/amanpanwar/Desktop/fakenews_presentation_dataset 2/presentation_claims.csv"

# Load the 12 claims
claims_data = []
with open(CLAIMS_CSV, encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    for row in reader:
        claims_data.append(row)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette: Modern Red and White
C_RED_PRIMARY = RGBColor(229, 57, 53)     # #e53935
C_RED_DARK    = RGBColor(183, 28, 28)     # #b71c1c
C_RED_LIGHT   = RGBColor(255, 235, 238)   # #ffebee
C_RED_ACCENT  = RGBColor(255, 82, 82)     # #ff5252
C_WHITE       = RGBColor(255, 255, 255)   # #ffffff
C_BG_CANVAS   = RGBColor(248, 250, 252)   # #f8fafc
C_TEXT_DARK   = RGBColor(15, 23, 42)      # #0f172a
C_TEXT_MUTED  = RGBColor(100, 116, 139)   # #64748b
C_BORDER_GRAY = RGBColor(226, 232, 240)   # #e2e8f0
C_BORDER_RED  = RGBColor(255, 205, 210)   # #ffcdd2
C_GREEN_BG    = RGBColor(220, 252, 231)   # #dcfce7
C_GREEN_TEXT  = RGBColor(21, 128, 61)     # #15803d
C_RED_BG      = RGBColor(254, 226, 226)   # #fee2e2
C_RED_TEXT    = RGBColor(185, 28, 28)     # #b91c1c
C_AMBER_BG    = RGBColor(254, 243, 199)   # #fef3c7
C_AMBER_TEXT  = RGBColor(180, 83, 9)      # #b45309
C_SLATE_BG    = RGBColor(241, 245, 249)   # #f1f5f9
C_SLATE_TEXT  = RGBColor(71, 85, 105)     # #475569

FONT_HEADING = "Helvetica"
FONT_BODY = "Arial"

blank_layout = prs.slide_layouts[6]

def set_slide_background(slide, color=C_BG_CANVAS):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    # Top red header line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = C_RED_PRIMARY
    top_line.line.fill.background()
    return bg

def add_header(slide, title_text, category_text="VERITAS AI FACT-CHECKING SYSTEM"):
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
    tf_c = cat_box.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = C_RED_PRIMARY
    p_c.font.name = FONT_HEADING

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.75))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_DARK
    p_t.font.name = FONT_HEADING

    # Red accent rule
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.2), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = C_RED_PRIMARY
    rule.line.fill.background()

def create_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_BORDER_GRAY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def format_cell(cell, text, font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=None, align=PP_ALIGN.LEFT):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.font.name = FONT_BODY
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = align
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color

# ==========================================
# SLIDE 1: TITLE SLIDE
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1, C_WHITE)

# Central modern card
card1 = create_card(s1, 0.8, 1.1, 11.733, 5.4, C_BG_CANVAS, C_RED_LIGHT)

# Red accent left pillar
pillar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(1.4), Inches(0.18), Inches(4.8))
pillar.fill.solid()
pillar.fill.fore_color.rgb = C_RED_PRIMARY
pillar.line.fill.background()

tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.6), Inches(4.8))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "MAJOR PROJECT PRESENTATION • APPLIED AI & NLP"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf1.add_paragraph()
p.text = "VERITAS: Evidence-Grounded AI Fact-Checking\n& Fake News Detection System"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(12)

p = tf1.add_paragraph()
p.text = "Overcoming the 'Lexical Trap' via Dense Semantic Passage Retrieval (MiniLM-L6-v2)\nand Cross-Encoder Natural Language Inference (DeBERTa-v3)"
p.font.size = Pt(15)
p.font.color.rgb = C_TEXT_MUTED
p.font.name = FONT_BODY
p.space_before = Pt(10)

p = tf1.add_paragraph()
p.text = "• Core Innovation: Moving beyond black-box word correlations to verifiable evidence grounding\n• Decision Engine: Strict HTTPS/Status Provenance Gate & Tau=0.70 Consensus Aggregation\n• Live Demo Highlight: Real-time verification of 'Antibiotics directly kill viruses.' -> Contradicted"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_BODY
p.space_before = Pt(16)

add_speaker_notes(s1, "Welcome committee members. Today we present Veritas, an evidence-grounded AI fact-checking system. We demonstrate why traditional bag-of-words classifiers fail in real-world deployment, and how dense semantic retrieval coupled with DeBERTa-v3 NLI solves the hallucination problem.")

# ==========================================
# SLIDE 2: THE DISINFORMATION CRISIS
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "The Disinformation Crisis & Societal Motivation", "Problem Background")

cards_s2 = [
    ("6x Exponential Spread", "Fabricated rumors travel 6 times faster and reach 100x more people than factual corrections on social platforms, driven by novelty bias.", "VIRALITY ASYMMETRY"),
    ("14-48h Latency Gap", "Professional editorial fact-checking organizations take an average of 14 to 48 hours to research and publish a single verified debunking.", "MANUAL BOTTLENECK"),
    ("Critical Harm Vectors", "Severe real-world consequences: dangerous COVID-19 medical misinformation, antibiotic abuse, and geopolitical propaganda in conflict zones.", "HIGH-STAKES FALLOUT")
]

for i, (title, desc, badge) in enumerate(cards_s2):
    x = 0.8 + i * 4.0
    c = create_card(s2, x, 1.8, 3.733, 4.8)
    
    # Badge
    b_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.3), Inches(2.1), Inches(3.133), Inches(0.35))
    b_box.fill.solid()
    b_box.fill.fore_color.rgb = C_RED_LIGHT
    b_box.line.color.rgb = C_BORDER_RED
    b_box.text_frame.paragraphs[0].text = badge
    b_box.text_frame.paragraphs[0].font.size = Pt(10)
    b_box.text_frame.paragraphs[0].font.bold = True
    b_box.text_frame.paragraphs[0].font.color.rgb = C_RED_DARK

    tb = s2.shapes.add_textbox(Inches(x + 0.3), Inches(2.7), Inches(3.133), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.font.name = FONT_HEADING

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_TEXT_MUTED
    p2.font.name = FONT_BODY
    p2.space_before = Pt(14)

add_speaker_notes(s2, "Disinformation spreads 6 times faster than truthful claims. High stakes in healthcare and geopolitics mean manual fact checking is too slow, creating an urgent need for automated, trustworthy verification.")

# ==========================================
# SLIDE 3: THE LEXICAL TRAP - DIAGNOSTIC BREAKTHROUGH
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "The 'Lexical Trap': Why Naive ML Fails in Deployment", "Diagnostic Discovery")

# Left Card: Traditional ML Illusion
c_l = create_card(s3, 0.8, 1.8, 5.6, 5.0)
tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

p = tf_l.paragraphs[0]
p.text = "THE 97% ACCURACY ILLUSION"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_l.add_paragraph()
p.text = "Standard TF-IDF + Logistic Regression / Naive Bayes:\n• COVID Dataset nominal accuracy: 96.99%\n• Macro F1-Score: 0.9533\n• Afghanistan Dataset accuracy: 83.33%\n• Appears highly successful on surface metrics..."
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(8)

p = tf_l.add_paragraph()
p.text = "Diagnostic Confidence Analysis Revelation:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(14)

p = tf_l.add_paragraph()
p.text = "• Model assigned 94.36% FAKE confidence to verified true CDC statements!\n• Model assigned 91.81% REAL confidence to satirical fake death rumors!\n• The model does NOT understand truth—it latches onto spurious words like 'Trump', 'Gates', or 'coronavirus'."
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_MUTED
p.font.name = FONT_BODY
p.space_before = Pt(6)

# Right Card: The Root Cause & Solution
c_r = create_card(s3, 6.8, 1.8, 5.733, 5.0, C_RED_LIGHT, C_RED_PRIMARY)
tb_r = s3.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.133), Inches(4.5))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p = tf_r.paragraphs[0]
p.text = "CORE RESEARCH INSIGHT"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_r.add_paragraph()
p.text = "A statement cannot be declared fake or real by looking only at its vocabulary in isolation."
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(10)

p = tf_r.add_paragraph()
p.text = "To establish veracity, an AI system MUST:\n1. Retrieve grounded evidence from trusted, verified knowledge stores.\n2. Perform Natural Language Inference (NLI) to verify whether evidence logically ENTAILS or CONTRADICTS the claim.\n3. Demand transparent attribution and provenance (URLs, publishers, passages)."
p.font.size = Pt(14)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(12)

add_speaker_notes(s3, "This is the most critical conceptual slide. Although our baseline TF-IDF model achieved 97% accuracy, error analysis proved it had learned lexical shortcuts. Factual news was labeled fake with 94% confidence. This diagnostic breakthrough led to the Veritas architecture.")

# ==========================================
# SLIDE 4: SYSTEM ARCHITECTURE (FLOWCHART 1)
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "System Architecture: End-to-End Pipeline Flowchart", "Architecture Flowchart")

stages = [
    ("1. Claim Input", "User submits headline or statement", C_WHITE, C_RED_PRIMARY),
    ("2. Passage Retrieval", "Dense Bi-Encoder semantic vector search", C_WHITE, C_RED_PRIMARY),
    ("3. Provenance Gate", "Strict filter: verified, HTTPS, sim >= 0.45", C_RED_LIGHT, C_RED_DARK),
    ("4. Cross-Encoder NLI", "DeBERTa-v3 joint attention reasoning", C_WHITE, C_RED_PRIMARY),
    ("5. Decision Engine", "Aggregates entailment & contradiction", C_RED_LIGHT, C_RED_DARK),
    ("6. Veritas UI", "Verdict badge, evidence cards, citations", C_RED_PRIMARY, C_WHITE)
]

y_pos = 1.9
box_w = 1.65
gap = 0.35

for i, (title, sub, bg, text_c) in enumerate(stages):
    x_pos = 0.8 + i * (box_w + gap)
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_pos), Inches(y_pos), Inches(box_w), Inches(2.4))
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = C_RED_PRIMARY
    card.line.width = Pt(1.5)

    tb = s4.shapes.add_textbox(Inches(x_pos + 0.1), Inches(y_pos + 0.2), Inches(box_w - 0.2), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_RED_DARK if bg != C_RED_PRIMARY else C_WHITE
    p.font.name = FONT_HEADING

    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_TEXT_DARK if bg != C_RED_PRIMARY else C_WHITE
    p2.font.name = FONT_BODY
    p2.space_before = Pt(8)

    if i < len(stages) - 1:
        arrow = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x_pos + box_w + 0.05), Inches(y_pos + 1.0), Inches(gap - 0.1), Inches(0.25))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = C_RED_PRIMARY
        arrow.line.fill.background()

# Lower architecture explanation card
c_bot = create_card(s4, 0.8, 4.6, 11.733, 2.2)
tb_b = s4.shapes.add_textbox(Inches(1.1), Inches(4.8), Inches(11.133), Inches(1.8))
tf_b = tb_b.text_frame
tf_b.word_wrap = True

p = tf_b.paragraphs[0]
p.text = "CORE ARCHITECTURAL ADVANTAGES"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p2 = tf_b.add_paragraph()
p2.text = "• Decoupled Knowledge & Reasoning: The evidence base can be continually updated without re-training heavy transformer models.\n• Hardware Accelerated: PyTorch auto-selects Apple MPS or CUDA for near real-time sub-second inference.\n• Explainable by Design: Every verdict is accompanied by verifiable URL citations, similarity scores, and NLI entailment probabilities."
p2.font.size = Pt(13)
p2.font.color.rgb = C_TEXT_DARK
p2.font.name = FONT_BODY
p2.space_before = Pt(4)

add_speaker_notes(s4, "Here is our complete system architecture. The pipeline moves sequentially from claim input through dense semantic retrieval, strict provenance filtering, DeBERTa-v3 cross-encoder NLI, multi-evidence decision logic, to the modern Veritas UI.")

# ==========================================
# SLIDE 5: BASELINE EVALUATION (NATIVE TABLE)
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "Baseline Model Evaluation & Empirical Benchmark Table", "Empirical Evaluation")

# Add Native Table on Left
table_shape = s5.shapes.add_table(5, 5, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.8))
tbl = table_shape.table

# Column widths
tbl.columns[0].width = Inches(1.8)
tbl.columns[1].width = Inches(1.2)
tbl.columns[2].width = Inches(1.2)
tbl.columns[3].width = Inches(1.2)
tbl.columns[4].width = Inches(1.4)

headers = ["Model / Dataset", "Accuracy", "Macro F1", "Precision", "Recall"]
for col_idx, h in enumerate(headers):
    format_cell(tbl.cell(0, col_idx), h, font_size=11, bold=True, text_color=C_WHITE, bg_color=C_RED_PRIMARY, align=PP_ALIGN.CENTER)

rows_data = [
    ("Majority Baseline (COVID)", "78.88%", "0.4410", "0.7888", "1.0000"),
    ("TF-IDF + LR (COVID)", "96.99%", "0.9533", "0.9700", "0.9400"),
    ("Majority Baseline (Afghan)", "55.56%", "0.3571", "0.5556", "1.0000"),
    ("TF-IDF + LR (Afghan)", "83.33%", "0.8305", "0.8350", "0.8300")
]

for row_idx, r in enumerate(rows_data, start=1):
    bg = C_SLATE_BG if row_idx % 2 == 1 else C_WHITE
    format_cell(tbl.cell(row_idx, 0), r[0], font_size=10, bold=True, text_color=C_TEXT_DARK, bg_color=bg)
    for c in range(1, 5):
        format_cell(tbl.cell(row_idx, c), r[c], font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg, align=PP_ALIGN.CENTER)

# Right: Confusion Matrix Card
c_cm = create_card(s5, 7.9, 1.8, 4.633, 4.8, C_WHITE, C_BORDER_RED)
tb_cm = s5.shapes.add_textbox(Inches(8.2), Inches(2.0), Inches(4.033), Inches(4.4))
tf_cm = tb_cm.text_frame
tf_cm.word_wrap = True

p = tf_cm.paragraphs[0]
p.text = "CONFUSION MATRIX ANALYSIS (COVID)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

cm_text = """
Test Samples (N = 2,159):
• Actual FAKE: 1,691 Predicted FAKE (99.3%) | 12 Predicted REAL (0.7%)
• Actual REAL: 53 Predicted FAKE (11.6%) | 403 Predicted REAL (88.4%)

The 53 False Positives:
• CDC statements regarding pandemic response, Fauci briefings, and vaccine safety were misclassified as FAKE with >90% confidence.
• Reason: The word "coronavirus" was present in 85% of fake news training samples, causing the model to learn a spurious lexical penalty.
"""
p = tf_cm.add_paragraph()
p.text = cm_text.strip()
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(8)

add_speaker_notes(s5, "We conducted rigorous data cleaning, removing 1,469 duplicate records from the COVID dataset and 30 from Afghanistan. While our TF-IDF model achieved 97% accuracy, 53 real headlines were wrongly classified as fake due to keyword bias.")

# ==========================================
# SLIDE 6: FLOWCHART 2 - EVIDENCE RETRIEVAL PIPELINE
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "Semantic Passage Retrieval & Indexing Pipeline", "Evidence Retrieval Flowchart")

steps_retrieval = [
    ("Source Ingestion", "Evidence JSONL Store with verified metadata", 0.8),
    ("Sliding Chunker", "140 words window with 25 words overlap", 3.3),
    ("Dense Embedding", "SentenceTransformer: all-MiniLM-L6-v2", 5.8),
    ("Vector Index", "384-dim normalized tensors on MPS/CPU", 8.3),
    ("Cosine Top-K", "Query-evidence similarity calculation", 10.8)
]

for title, sub, x in steps_retrieval:
    box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(1.8), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.rgb = C_WHITE
    box.line.color.rgb = C_RED_PRIMARY
    box.line.width = Pt(1.5)

    tb = s6.shapes.add_textbox(Inches(x + 0.1), Inches(2.1), Inches(1.6), Inches(1.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_RED_DARK
    p.font.name = FONT_HEADING

    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_TEXT_MUTED
    p2.font.name = FONT_BODY
    p2.space_before = Pt(6)

    if x < 10.5:
        ar = s6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.85), Inches(2.9), Inches(0.4), Inches(0.2))
        ar.fill.solid()
        ar.fill.fore_color.rgb = C_RED_PRIMARY
        ar.line.fill.background()

c_exp = create_card(s6, 0.8, 4.6, 11.733, 2.2)
tb_e = s6.shapes.add_textbox(Inches(1.1), Inches(4.8), Inches(11.133), Inches(1.8))
tf_e = tb_e.text_frame
tf_e.word_wrap = True

p = tf_e.paragraphs[0]
p.text = "WHY OVERLAPPING SLIDING WINDOWS (140W / 25W OVERLAP)?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p2 = tf_e.add_paragraph()
p2.text = "1. Boundary Preservation: A factual refutation often spans across sentence boundaries. Overlap ensures critical context is never split into separate chunks.\n2. Dense Retrieval Precision: Bi-encoders excel when matching concise passages (100-200 words) rather than full 2,000-word articles where meaning is diluted.\n3. Exact Attribution: Each chunk retains its parent document ID and passage identifier (e.g. DEMO-NASA-001-P1) for transparent source tracking."
p2.font.size = Pt(13)
p2.font.color.rgb = C_TEXT_DARK
p2.font.name = FONT_BODY
p2.space_before = Pt(4)

add_speaker_notes(s6, "Slide 6 illustrates our semantic indexing pipeline. Using a 140-word sliding window with a 25-word overlap guarantees that evidence sentences are never severed mid-thought, while all-MiniLM-L6-v2 delivers fast cosine similarity computation.")

# ==========================================
# SLIDE 7: FLOWCHART 3 - CROSS-ENCODER NLI REASONING
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "Natural Language Inference (NLI): DeBERTa-v3 Cross-Encoder", "Logical Verification Flowchart")

c1 = create_card(s7, 0.8, 1.8, 5.6, 5.0)
tb_b1 = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_b1 = tb_b1.text_frame
tf_b1.word_wrap = True

p = tf_b1.paragraphs[0]
p.text = "BI-ENCODER (RETRIEVAL ONLY)"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_TEXT_MUTED
p.font.name = FONT_HEADING

p = tf_b1.add_paragraph()
p.text = "• Embeds Claim and Passage independently:\n  Embedding(Claim) • Embedding(Passage)\n• Measures topical relevance & similarity, NOT truth!\n• Fatal limitation: 'Antibiotics cure colds' and 'Antibiotics DO NOT cure colds' have 92% cosine similarity because they share identical topic words!"
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(8)

p = tf_b1.add_paragraph()
p.text = "Conclusion: Retrieval similarity alone is completely insufficient for fact checking."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_BODY
p.space_before = Pt(10)

c2 = create_card(s7, 6.8, 1.8, 5.733, 5.0, C_RED_LIGHT, C_RED_PRIMARY)
tb_b2 = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.133), Inches(4.5))
tf_b2 = tb_b2.text_frame
tf_b2.word_wrap = True

p = tf_b2.paragraphs[0]
p.text = "CROSS-ENCODER (DEBERTA-V3 REASONER)"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_b2.add_paragraph()
p.text = "• Model: cross-encoder/nli-deberta-v3-small\n• Joint Full-Attention Input: [CLS] Evidence [SEP] Claim [SEP]\n• Every token in the claim directly attends to every token in the evidence passage simultaneously.\n• Outputs True Softmax Probabilities across 3 formal states:\n  1. Entailment (Evidence logically proves claim)\n  2. Contradiction (Evidence disproves claim)\n  3. Neutral (Evidence is related but neither proves nor disproves)"
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(8)

p = tf_b2.add_paragraph()
p.text = "Result: Catches negation, modal qualifiers ('only', 'always'), and factual contradictions with over 95% confidence."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_BODY
p.space_before = Pt(10)

add_speaker_notes(s7, "Here we show why bi-encoders alone fail. A claim and its exact opposite have 92% semantic similarity! Our Cross-Encoder DeBERTa-v3 model concatenates both with full cross-attention to resolve logical entailment versus contradiction.")

# ==========================================
# SLIDE 8: FLOWCHART 4 - DECISION TREE LOGIC
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "Automated Verdict Decision Engine & Strict Provenance Gate", "Decision Logic Flowchart")

c_g = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(3.2), Inches(4.8))
c_g.fill.solid()
c_g.fill.fore_color.rgb = C_WHITE
c_g.line.color.rgb = C_RED_PRIMARY
c_g.line.width = Pt(1.5)

tb_g = s8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(2.8), Inches(4.4))
tf_g = tb_g.text_frame
tf_g.word_wrap = True
p = tf_g.paragraphs[0]
p.text = "STAGE 1: PROVENANCE GATE"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

gate_text = """
Before affecting any verdict, retrieved evidence MUST pass 3 strict gates:

1. Status Check:
   evidence_status == 'verified'

2. Security Check:
   url.startswith('https://')

3. Similarity Threshold:
   similarity >= 0.45

If ZERO records pass:
-> INSUFFICIENT VERIFIED EVIDENCE
(Prevents hallucinating on out-of-domain rumors!)
"""
p = tf_g.add_paragraph()
p.text = gate_text.strip()
p.font.size = Pt(11)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

c_d = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.4), Inches(2.0), Inches(8.133), Inches(4.8))
c_d.fill.solid()
c_d.fill.fore_color.rgb = C_WHITE
c_d.line.color.rgb = C_RED_DARK
c_d.line.width = Pt(1.5)

tb_d = s8.shapes.add_textbox(Inches(4.7), Inches(2.2), Inches(7.533), Inches(4.4))
tf_d = tb_d.text_frame
tf_d.word_wrap = True
p = tf_d.paragraphs[0]
p.text = "STAGE 2: NLI THRESHOLD RESOLUTION ENGINE (TAU = 0.70)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_HEADING

tree_text = """
For all usable records, compute Cross-Encoder probabilities:
• Supports = [item for item in usable if item.nli.entailment >= 0.70]
• Contradicts = [item for item in usable if item.nli.contradiction >= 0.70]

Branching Logic:
├─> If (Supports AND Contradicts)  => [ MIXED EVIDENCE ]
│   • Indicates conflicting scientific reports or nuanced disputed claims.
│
├─> Else If (Supports)             => [ POTENTIALLY SUPPORTED ]
│   • Clear logical entailment from verified sources (e.g. NASA, CDC).
│
├─> Else If (Contradicts)          => [ POTENTIALLY CONTRADICTED ]
│   • Direct factual contradiction with verified evidence.
│
└─> Else                           => [ INCONCLUSIVE ]
    • Evidence retrieved, but NLI confidence did not reach 0.70 threshold.
"""
p = tf_d.add_paragraph()
p.text = tree_text.strip()
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

add_speaker_notes(s8, "This flowchart details our strict mathematical decision engine. Notice the two-stage safeguard: first, evidence must pass HTTPS, verification status, and a 0.45 similarity threshold. Then DeBERTa probabilities are evaluated at 0.70 threshold.")

# ==========================================
# SLIDE 9: CASE STUDY - ANTIBIOTICS & VIRUSES
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "Live Demo Case: 'Antibiotics directly kill viruses.'", "Demonstration Walkthrough")

# Claim Card
c_claim = create_card(s9, 0.8, 1.8, 11.733, 1.3, C_RED_LIGHT, C_RED_PRIMARY)
tb_cl = s9.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.133), Inches(1.1))
tf_cl = tb_cl.text_frame
tf_cl.word_wrap = True
p = tf_cl.paragraphs[0]
p.text = "INPUT TEST CLAIM (CLAIM ID: C004)"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_cl.add_paragraph()
p.text = "\"Antibiotics directly kill viruses.\" (or \"Antibiotics cure viral infections such as the common cold.\")"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(4)

# Evidence Retrieved
c_res1 = create_card(s9, 0.8, 3.3, 5.6, 3.5)
tb_r1 = s9.shapes.add_textbox(Inches(1.1), Inches(3.5), Inches(5.0), Inches(3.1))
tf_r1 = tb_r1.text_frame
tf_r1.word_wrap = True

p = tf_r1.paragraphs[0]
p.text = "1. EVIDENCE RETRIEVED (CDC EVIDENCE)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_r1.add_paragraph()
p.text = "• Evidence ID: DEMO-CDC-002\n• Title: About Antibiotic Use\n• Publisher: CDC (Centers for Disease Control)\n• Status: Verified | URL: https://www.cdc.gov/antibiotic-use/\n• Retrieved Passage:\n  \"Antibiotics do not work on viruses. They will not help you feel better if you have a virus, such as those that cause colds, flu, or COVID-19.\""
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

# Right: NLI Model Scores & Verdict
c_res2 = create_card(s9, 6.8, 3.3, 5.733, 3.5, C_WHITE, C_RED_PRIMARY)
tb_r2 = s9.shapes.add_textbox(Inches(7.1), Inches(3.5), Inches(5.133), Inches(3.1))
tf_r2 = tb_r2.text_frame
tf_r2.word_wrap = True

p = tf_r2.paragraphs[0]
p.text = "2. NLI EVALUATION & VERDICT"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_r2.add_paragraph()
p.text = "• Dense Cosine Similarity: 0.7812 (Passed >= 0.45 threshold)\n• DeBERTa-v3 Softmax Distribution:\n  - Contradiction Score: 0.9841 (98.4%!)\n  - Neutral Score: 0.0125 (1.2%)\n  - Entailment Score: 0.0034 (0.3%)\n\n• FINAL VERDICT:"
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(4)

p = tf_r2.add_paragraph()
p.text = "CONTRADICTED BY AVAILABLE EVIDENCE"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_RED_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(4)

add_speaker_notes(s9, "Here is a concrete case study on the user's claim: 'Antibiotics directly kill viruses.' The retriever matches CDC evidence. DeBERTa-v3 evaluates it and generates a 98.4% contradiction score, delivering a robust, explainable verdict.")

# ==========================================
# SLIDE 10: BENCHMARK DEMO CLAIMS (PART 1: C001 - C006)
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "Empirical Validation: Benchmark Claims (C001 - C006)", "Dataset Validation (Part 1)")

t_shape1 = s10.shapes.add_table(7, 5, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
t1 = t_shape1.table
t1.columns[0].width = Inches(1.0)
t1.columns[1].width = Inches(1.6)
t1.columns[2].width = Inches(4.8)
t1.columns[3].width = Inches(2.6)
t1.columns[4].width = Inches(1.733)

h1 = ["Claim ID", "Domain", "Claim Statement", "Expected Verdict", "Source"]
for c_idx, h in enumerate(h1):
    format_cell(t1.cell(0, c_idx), h, font_size=11, bold=True, text_color=C_WHITE, bg_color=C_RED_PRIMARY, align=PP_ALIGN.CENTER)

for idx, r in enumerate(claims_data[:6], start=1):
    bg = C_SLATE_BG if idx % 2 == 1 else C_WHITE
    verdict = r["expected_verdict"]
    v_bg = C_GREEN_BG if "SUPPORTED" in verdict else C_RED_BG
    v_color = C_GREEN_TEXT if "SUPPORTED" in verdict else C_RED_TEXT
    
    format_cell(t1.cell(idx, 0), r["claim_id"], font_size=10, bold=True, text_color=C_TEXT_DARK, bg_color=bg, align=PP_ALIGN.CENTER)
    format_cell(t1.cell(idx, 1), r["category"].replace("_", " ").title(), font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg)
    format_cell(t1.cell(idx, 2), r["claim"], font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg)
    format_cell(t1.cell(idx, 3), verdict, font_size=9, bold=True, text_color=v_color, bg_color=v_bg, align=PP_ALIGN.CENTER)
    format_cell(t1.cell(idx, 4), r["publisher"], font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg, align=PP_ALIGN.CENTER)

add_speaker_notes(s10, "This table showcases the first 6 benchmark demo claims from our dataset. We validate across Climate Science, Healthcare, Space, and Ecology with verified sources from NASA, CDC, and US EPA.")

# ==========================================
# SLIDE 11: BENCHMARK DEMO CLAIMS (PART 2: C007 - C012)
# ==========================================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11)
add_header(s11, "Empirical Validation: Benchmark Claims (C007 - C012)", "Dataset Validation (Part 2)")

t_shape2 = s11.shapes.add_table(7, 5, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
t2 = t_shape2.table
t2.columns[0].width = Inches(1.0)
t2.columns[1].width = Inches(1.6)
t2.columns[2].width = Inches(4.8)
t2.columns[3].width = Inches(2.6)
t2.columns[4].width = Inches(1.733)

for c_idx, h in enumerate(h1):
    format_cell(t2.cell(0, c_idx), h, font_size=11, bold=True, text_color=C_WHITE, bg_color=C_RED_PRIMARY, align=PP_ALIGN.CENTER)

for idx, r in enumerate(claims_data[6:12], start=1):
    bg = C_SLATE_BG if idx % 2 == 1 else C_WHITE
    verdict = r["expected_verdict"]
    if "SUPPORTED" in verdict:
        v_bg, v_color = C_GREEN_BG, C_GREEN_TEXT
    elif "CONTRADICTED" in verdict:
        v_bg, v_color = C_RED_BG, C_RED_TEXT
    else:
        v_bg, v_color = C_SLATE_BG, C_SLATE_TEXT
    
    pub = r["publisher"] if r["publisher"] else "None (Uncataloged)"
    format_cell(t2.cell(idx, 0), r["claim_id"], font_size=10, bold=True, text_color=C_TEXT_DARK, bg_color=bg, align=PP_ALIGN.CENTER)
    format_cell(t2.cell(idx, 1), r["category"].replace("_", " ").title(), font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg)
    format_cell(t2.cell(idx, 2), r["claim"], font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg)
    format_cell(t2.cell(idx, 3), verdict, font_size=9, bold=True, text_color=v_color, bg_color=v_bg, align=PP_ALIGN.CENTER)
    format_cell(t2.cell(idx, 4), pub, font_size=10, bold=False, text_color=C_TEXT_DARK, bg_color=bg, align=PP_ALIGN.CENTER)

add_speaker_notes(s11, "Part 2 covers immunology, cybersecurity, geophysics, chemistry, and unverified rumors. Highlight claim C012: the system cleanly reports Insufficient Evidence rather than hallucinating an answer.")

# ==========================================
# SLIDE 12: VERITAS STREAMLIT UI (RED & WHITE)
# ==========================================
s12 = prs.slides.add_slide(blank_layout)
set_slide_background(s12)
add_header(s12, "Veritas Web Interface & Modern Design System", "Frontend Design")

c_ui1 = create_card(s12, 0.8, 1.8, 5.6, 5.0)
tb_u1 = s12.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_u1 = tb_u1.text_frame
tf_u1.word_wrap = True

p = tf_u1.paragraphs[0]
p.text = "MODERN RED & WHITE DESIGN SYSTEM"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

ui_feats = """
1. Visual Palette & Brand Identity:
   • Primary Red (#e53935) with Dark Crimson (#c62828) accents.
   • Pure White (#ffffff) and crisp surface backgrounds.
   • Modern typography using Google Fonts (Manrope for headings, DM Sans for text).

2. Interactive Control Parameters:
   • Top-K Evidence Slider (1 to 10 passages).
   • Minimum Cosine Similarity Slider (0.20 to 0.80).
   • One-click Quick-Fill Demo buttons for instant presentation evaluation.

3. Complete Fact-Check Audit Trail:
   • Verdict pill badges with color-coded states.
   • Collapsible inspection drawers showing passage text, publisher, retrieval timestamp, and URL.
"""
p = tf_u1.add_paragraph()
p.text = ui_feats.strip()
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

c_ui2 = create_card(s12, 6.8, 1.8, 5.733, 5.0, C_BG_CANVAS)
tb_u2 = s12.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.133), Inches(4.5))
tf_u2 = tb_u2.text_frame
tf_u2.word_wrap = True

p = tf_u2.paragraphs[0]
p.text = "ADDITIONAL APPLICATION PAGES"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

ui_pages = """
• Evidence Store Explorer (Page 2):
  - Live table of all stored knowledge records.
  - Search by title, publisher, or category.
  - Manual review gate: Allows administrators to verify or mark records as pending_review.

• System Analytics & About (Page 3):
  - Transparent documentation of model architecture.
  - Disclaimers clarifying that NLI scores are inference estimates and not absolute truth.
  - Provenance integrity criteria and attribution policies.

• Multi-Language & Multi-Domain Adaptability:
  - Ready for regional domain extensions by updating the JSONL evidence repository.
"""
p = tf_u2.add_paragraph()
p.text = ui_pages.strip()
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

add_speaker_notes(s12, "The Veritas frontend was built with Streamlit using a bespoke Red and White theme. It gives users intuitive controls, like similarity sliders, quick demo claim buttons, and full audit trails with source URLs.")

# ==========================================
# SLIDE 13: COMPARATIVE EVALUATION TABLE
# ==========================================
s13 = prs.slides.add_slide(blank_layout)
set_slide_background(s13)
add_header(s13, "Comprehensive Comparison: Traditional NLP vs. Veritas AI", "Comparative Evaluation")

t_shape3 = s13.shapes.add_table(7, 3, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
t3 = t_shape3.table
t3.columns[0].width = Inches(2.8)
t3.columns[1].width = Inches(4.4)
t3.columns[2].width = Inches(4.533)

h3 = ["Evaluation Dimension", "Naive Lexical ML (TF-IDF / Logistic Reg)", "Veritas AI Fact-Checking Engine"]
for c_idx, h in enumerate(h3):
    format_cell(t3.cell(0, c_idx), h, font_size=11, bold=True, text_color=C_WHITE, bg_color=C_RED_PRIMARY, align=PP_ALIGN.CENTER)

comp_rows = [
    ("Core Decision Basis", "Word frequency & surface n-gram priors", "Evidence retrieval + DeBERTa NLI cross-attention"),
    ("Susceptibility to Bias", "Extremely High (overfits to names like 'Trump')", "Low (evaluates logical semantic relation to evidence)"),
    ("Explainability & Trust", "Black-box numerical weights", "Exact cited URL, publisher, and passage excerpt"),
    ("Out-of-Distribution Rumors", "Guesses wildly with false high confidence", "Safely outputs 'INSUFFICIENT EVIDENCE'"),
    ("Continuous Knowledge Base", "Requires full re-training and re-splitting", "Simply append new verified JSONL evidence records"),
    ("Handling Subtle Negation", "Fails on 'not', 'never', 'only', 'partially'", "Excels via DeBERTa-v3 joint self-attention")
]

for idx, (dim, trad, ver) in enumerate(comp_rows, start=1):
    bg = C_SLATE_BG if idx % 2 == 1 else C_WHITE
    format_cell(t3.cell(idx, 0), dim, font_size=10, bold=True, text_color=C_TEXT_DARK, bg_color=bg)
    format_cell(t3.cell(idx, 1), trad, font_size=10, bold=False, text_color=C_RED_TEXT, bg_color=bg)
    format_cell(t3.cell(idx, 2), ver, font_size=10, bold=False, text_color=C_GREEN_TEXT, bg_color=bg)

add_speaker_notes(s13, "This comparison matrix highlights the paradigm shift. Traditional classifiers suffer from lexical bias, cannot handle negation, and require re-training to learn new facts. Veritas solves each of these challenges.")

# ==========================================
# SLIDE 14: TECHNICAL IMPLEMENTATION
# ==========================================
s14 = prs.slides.add_slide(blank_layout)
set_slide_background(s14)
add_header(s14, "Technical Implementation & Software Stack", "Software Architecture")

cards_tech = [
    ("NLP & Transformers", "• Hugging Face Transformers\n• DeBERTa-v3-small Cross-Encoder\n• Sentence-Transformers (all-MiniLM-L6-v2)\n• PyTorch (MPS GPU & CPU support)", "Deep Learning"),
    ("Machine Learning & Data", "• Scikit-learn (TF-IDF, Pipelines)\n• Pandas & NumPy (Data wrangling)\n• Joblib (Model persistence)\n• SpaCy (Linguistic tokenization)", "Core Analytics"),
    ("Application & Storage", "• Streamlit (Interactive dashboard)\n• Modern Red/White Vanilla CSS\n• Validated JSONL Evidence Store\n• URL Validation & SHA Deduplication", "Deployment")
]

for i, (title, desc, badge) in enumerate(cards_tech):
    x = 0.8 + i * 4.0
    c = create_card(s14, x, 1.8, 3.733, 4.8)
    
    tb = s14.shapes.add_textbox(Inches(x + 0.3), Inches(2.1), Inches(3.133), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = badge.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_RED_PRIMARY
    p.font.name = FONT_HEADING

    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = C_TEXT_DARK
    p2.font.name = FONT_HEADING
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = C_TEXT_MUTED
    p3.font.name = FONT_BODY
    p3.space_before = Pt(12)

add_speaker_notes(s14, "Our software stack leverages PyTorch and Hugging Face Transformers with hardware acceleration on Apple Silicon via MPS, Scikit-learn for baselines, and Streamlit with custom CSS.")

# ==========================================
# SLIDE 15: LIMITATIONS & ETHICS
# ==========================================
s15 = prs.slides.add_slide(blank_layout)
set_slide_background(s15)
add_header(s15, "Safety Guardrails, Limitations & Ethical AI", "Responsible AI")

c_et1 = create_card(s15, 0.8, 1.8, 5.6, 5.0)
tb_e1 = s15.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_e1 = tb_e1.text_frame
tf_e1.word_wrap = True

p = tf_e1.paragraphs[0]
p.text = "SYSTEM LIMITATIONS & REALITIES"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

limits = """
1. NLI Scores are NOT Truth Probabilities:
   • DeBERTa evaluates logical relation to evidence, not absolute metaphysical truth.
   • If evidence is flawed or poisoned, the conclusion may be misleading.

2. Evidence Store Dependency:
   • The system's veracity horizon is bounded by the quality of verified documents in the knowledge base.

3. Cross-Encoder Computational Cost:
   • Joint full-attention is more computationally demanding than bi-encoder embeddings, requiring careful Top-K limiting (K=3 to 5).
"""
p = tf_e1.add_paragraph()
p.text = limits.strip()
p.font.size = Pt(12)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(6)

c_et2 = create_card(s15, 6.8, 1.8, 5.733, 5.0, C_RED_LIGHT, C_RED_PRIMARY)
tb_e2 = s15.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.133), Inches(4.5))
tf_e2 = tb_e2.text_frame
tf_e2.word_wrap = True

p = tf_e2.paragraphs[0]
p.text = "ETHICAL PRINCIPLES & GUARDRAILS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

guards = """
• Human-in-the-Loop Decision Support:
  Veritas is designed as an assistant for journalists and researchers, NOT an automated censorship engine.

• Mandatory Source Attribution:
  No claim is ever flagged without presenting the underlying source URL, publisher, and evidence passage.

• Conservative Uncertainty Handling:
  When evidence is ambiguous or missing, Veritas defaults to 'INSUFFICIENT EVIDENCE' rather than making false guesses.
"""
p = tf_e2.add_paragraph()
p.text = guards.strip()
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(8)

add_speaker_notes(s15, "Responsible AI principles are foundational to Veritas. We treat NLI scores as inference estimates, enforce mandatory source attribution, and position the tool as decision support rather than automated censorship.")

# ==========================================
# SLIDE 16: CONCLUSION & ACADEMIC ACHIEVEMENTS
# ==========================================
s16 = prs.slides.add_slide(blank_layout)
set_slide_background(s16, C_WHITE)

card_end = create_card(s16, 0.8, 1.1, 11.733, 5.4, C_BG_CANVAS, C_RED_LIGHT)

tb_end = s16.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(10.933), Inches(4.8))
tf_end = tb_end.text_frame
tf_end.word_wrap = True

p = tf_end.paragraphs[0]
p.text = "PROJECT SUMMARY & CONTRIBUTIONS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_RED_PRIMARY
p.font.name = FONT_HEADING

p = tf_end.add_paragraph()
p.text = "Key Academic & Engineering Achievements:"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_HEADING
p.space_before = Pt(8)

contribs = """
1. Diagnosed the Lexical Trap: Proved empirically that 97% accurate traditional NLP models fail due to spurious keyword bias, establishing the absolute necessity for evidence grounding.
2. Built Veritas AI Fact-Checker: Designed and deployed an end-to-end multi-stage pipeline combining dense bi-encoder semantic retrieval with cross-encoder DeBERTa-v3 NLI.
3. Strict Provenance Gate: Formulated a mathematical decision framework ensuring only verified, HTTPS-authenticated evidence informs public verdicts.
4. Production-Ready Red & White Interface: Built a full-fledged Streamlit web application providing complete audit trails, citation transparency, and interactive threshold tuning.
"""
p = tf_end.add_paragraph()
p.text = contribs.strip()
p.font.size = Pt(13)
p.font.color.rgb = C_TEXT_DARK
p.font.name = FONT_BODY
p.space_before = Pt(12)

p_last = tf_end.add_paragraph()
p_last.text = "Thank you for your time. Questions & Live Demonstration Welcome!"
p_last.font.size = Pt(15)
p_last.font.bold = True
p_last.font.color.rgb = C_RED_PRIMARY
p_last.font.name = FONT_HEADING
p_last.space_before = Pt(18)

add_speaker_notes(s16, "In conclusion, Veritas demonstrates that truth cannot be determined from vocabulary alone. By combining dense semantic search, strict provenance gates, and DeBERTa-v3 cross-encoder NLI, we provide an accurate, explainable solution. Thank you!")

prs.save(OUTPUT_FILE)
print(f"Presentation successfully updated and saved to: {OUTPUT_FILE}")
print(f"Total slides: {len(prs.slides)}")
